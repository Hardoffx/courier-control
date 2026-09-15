from functools import wraps
from datetime import datetime
from io import BytesIO
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.db import transaction
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST
from openpyxl import load_workbook
from accounts.models import User
from .forms import DeliveryForm
from .models import Delivery, DeliveryEvent, DeliveryPoint

def dispatcher_required(view):
    @wraps(view)
    @login_required
    def wrapped(request,*args,**kwargs):
        if not request.user.is_dispatcher: raise PermissionDenied
        return view(request,*args,**kwargs)
    return wrapped

@login_required
def home(request): return redirect('dispatcher_dashboard' if request.user.is_dispatcher else 'courier_today')

@dispatcher_required
def dispatcher_dashboard(request):
    today=timezone.localdate(); deliveries=Delivery.objects.filter(delivery_date=today).select_related('courier','point'); couriers=User.objects.filter(role=User.Role.COURIER,is_active=True).order_by('first_name','username'); counts={key:deliveries.filter(status=key).count() for key,_ in Delivery.Status.choices}; courier_stats=[]
    for courier in couriers:
        qs=deliveries.filter(courier=courier); courier_stats.append({'courier':courier,'total':qs.count(),'done':qs.filter(status=Delivery.Status.DONE).count(),'problem':qs.filter(status=Delivery.Status.PROBLEM).count()})
    return render(request,'dispatcher/dashboard.html',{'deliveries':deliveries,'couriers':couriers,'row_colors':Delivery.RowColor.choices,'counts':counts,'courier_stats':courier_stats,'today':today})

def _resolve_courier(courier_id):
    if not courier_id: return None
    return get_object_or_404(User,pk=courier_id,role=User.Role.COURIER,is_active=True)

def _assign(delivery,courier,actor,action='assigned'):
    old=delivery.courier; delivery.courier=courier
    if courier and delivery.status==Delivery.Status.NEW: delivery.status=Delivery.Status.IN_PROGRESS
    elif not courier and delivery.status==Delivery.Status.IN_PROGRESS: delivery.status=Delivery.Status.NEW
    delivery.save(update_fields=['courier','status','updated_at']); DeliveryEvent.objects.create(delivery=delivery,actor=actor,action=action,note=f'{old or "—"} → {courier or "—"}')

@dispatcher_required
@require_POST
def dispatcher_assign(request,pk):
    _assign(get_object_or_404(Delivery,pk=pk),_resolve_courier(request.POST.get('courier_id','')),request.user); return redirect('dispatcher_dashboard')

@dispatcher_required
@require_POST
def dispatcher_bulk_assign(request):
    ids=request.POST.getlist('delivery_ids'); courier=_resolve_courier(request.POST.get('courier_id','')); deliveries=Delivery.objects.filter(pk__in=ids,delivery_date=timezone.localdate()).exclude(status=Delivery.Status.DONE)
    with transaction.atomic():
        for delivery in deliveries: _assign(delivery,courier,request.user,'bulk_assigned')
    messages.success(request,f'Обновлено точек: {deliveries.count()}') if ids else messages.warning(request,'Сначала отметьте точки'); return redirect('dispatcher_dashboard')

@dispatcher_required
@require_POST
def dispatcher_quick_edit(request,pk):
    delivery=get_object_or_404(Delivery,pk=pk,delivery_date=timezone.localdate()); changed=[]
    values={'source_label':request.POST.get('source_label','').strip()[:255],'address':request.POST.get('address','').strip()[:500],'time_window':request.POST.get('time_window','').strip()[:64],'phone':request.POST.get('phone','').strip()[:64],'row_color':request.POST.get('row_color','')}
    allowed_colors={v for v,_ in Delivery.RowColor.choices}
    if values['row_color'] not in allowed_colors: values['row_color']=''
    if not values['address']: messages.error(request,'Адрес не может быть пустым'); return redirect('dispatcher_dashboard')
    for field,value in values.items():
        if getattr(delivery,field)!=value: setattr(delivery,field,value); changed.append(field)
    if changed:
        delivery.save(update_fields=changed+['updated_at']); DeliveryEvent.objects.create(delivery=delivery,actor=request.user,action='quick_edited',note=', '.join(changed))
    return redirect('dispatcher_dashboard')

@dispatcher_required
def delivery_create(request):
    form=DeliveryForm(request.POST or None,initial={'delivery_date':timezone.localdate()})
    if request.method=='POST' and form.is_valid():
        delivery=form.save(); DeliveryEvent.objects.create(delivery=delivery,actor=request.user,action='created'); return redirect('dispatcher_dashboard')
    return render(request,'dispatcher/delivery_form.html',{'form':form,'title':'Новая заявка'})

@dispatcher_required
def delivery_edit(request,pk):
    delivery=get_object_or_404(Delivery,pk=pk); form=DeliveryForm(request.POST or None,instance=delivery)
    if request.method=='POST' and form.is_valid(): form.save(); DeliveryEvent.objects.create(delivery=delivery,actor=request.user,action='edited'); return redirect('dispatcher_dashboard')
    return render(request,'dispatcher/delivery_form.html',{'form':form,'title':'Редактирование заявки'})

@dispatcher_required
def import_excel(request):
    if request.method=='POST' and request.FILES.get('file'):
        try:
            wb=load_workbook(BytesIO(request.FILES['file'].read()),data_only=True); ws=wb.active; rows=list(ws.iter_rows(values_only=True)); headers=[str(v or '').strip().lower() for v in rows[0]] if rows else []; aliases={'адрес':'address','address':'address','объект':'source_label','код':'source_label','точка':'source_label','организация':'organization','получатель':'recipient','телефон':'phone','комментарий':'comment','курьер':'courier','порядок':'route_order','№':'route_order','дата':'delivery_date','время':'time_window','временное окно':'time_window'}; columns={aliases[h]:i for i,h in enumerate(headers) if h in aliases}
            if 'address' not in columns: raise ValueError('Не найдена колонка «Адрес»')
            created=0
            with transaction.atomic():
                for number,row in enumerate(rows[1:],start=1):
                    address=str(row[columns['address']] or '').strip()
                    if not address: continue
                    def value(key): return str(row[columns[key]] or '').strip() if key in columns else ''
                    label=value('source_label'); kind=Delivery.infer_point_kind(label); point=None
                    if label: point,_=DeliveryPoint.objects.get_or_create(code=label if kind==DeliveryPoint.Kind.CMD else '',address=address,defaults={'name':label,'kind':kind})
                    courier=User.objects.filter(role=User.Role.COURIER,username__iexact=value('courier')).first() if value('courier') else None; date=timezone.localdate()
                    if 'delivery_date' in columns and row[columns['delivery_date']]:
                        raw=row[columns['delivery_date']]; date=raw.date() if isinstance(raw,datetime) else raw if hasattr(raw,'year') else date
                    order=number
                    if 'route_order' in columns and row[columns['route_order']] is not None:
                        try: order=int(row[columns['route_order']])
                        except (TypeError,ValueError): pass
                    row_color=''
                    try: row_color={'FFFFFF00':'yellow','FFFFC000':'yellow','FF92D050':'green','FFC6E0B4':'green','FFFFC7CE':'red','FFF4CCCC':'red','FFD9EAF7':'blue','FFD9EAD3':'green','FFD9D9D9':'gray'}.get(getattr(ws.cell(row=number+1,column=columns['address']+1).fill.fgColor,'rgb',None),'')
                    except Exception: pass
                    d=Delivery.objects.create(delivery_date=date,point=point,source_label=label,address=address,organization=value('organization'),recipient=value('recipient'),phone=value('phone'),comment=value('comment'),time_window=value('time_window'),row_color=row_color,courier=courier,route_order=order,status=Delivery.Status.IN_PROGRESS if courier else Delivery.Status.NEW); DeliveryEvent.objects.create(delivery=d,actor=request.user,action='imported',note=f'Тип: {kind}'); created+=1
            messages.success(request,f'Импортировано заявок: {created}'); return redirect('dispatcher_dashboard')
        except Exception as exc: messages.error(request,f'Не удалось импортировать файл: {exc}')
    return render(request,'dispatcher/import_excel.html')

@login_required
def courier_today(request):
    if request.user.is_dispatcher: return redirect('dispatcher_dashboard')
    today=timezone.localdate(); deliveries=Delivery.objects.filter(delivery_date=today,courier=request.user).order_by('route_order','id'); done=deliveries.filter(status=Delivery.Status.DONE).count(); return render(request,'courier/today.html',{'deliveries':deliveries,'done':done,'total':deliveries.count(),'today':today})

@login_required
@require_POST
def courier_update(request,pk):
    delivery=get_object_or_404(Delivery,pk=pk,courier=request.user); action=request.POST.get('action')
    if action=='done': delivery.status=Delivery.Status.DONE; delivery.completed_at=timezone.now(); delivery.completed_latitude=request.POST.get('latitude') or None; delivery.completed_longitude=request.POST.get('longitude') or None
    elif action=='problem': delivery.status=Delivery.Status.PROBLEM; delivery.problem_reason=request.POST.get('problem_reason','Другая проблема')[:255]
    elif action=='phone': delivery.phone=request.POST.get('phone','')[:64]
    else: raise PermissionDenied
    delivery.save(); DeliveryEvent.objects.create(delivery=delivery,actor=request.user,action=action); return redirect('courier_today')

@login_required
@require_POST
def courier_reorder(request,pk):
    delivery=get_object_or_404(Delivery,pk=pk,courier=request.user,delivery_date=timezone.localdate()); direction=request.POST.get('direction'); items=list(Delivery.objects.filter(courier=request.user,delivery_date=delivery.delivery_date).exclude(status=Delivery.Status.DONE).order_by('route_order','id'))
    try: index=items.index(delivery)
    except ValueError: return redirect('courier_today')
    target=index-1 if direction=='up' else index+1
    if 0<=target<len(items):
        other=items[target]; a,b=delivery.route_order,other.route_order; delivery.route_order=b if a!=b else target+1; other.route_order=a if a!=b else index+1; delivery.save(update_fields=['route_order']); other.save(update_fields=['route_order']); DeliveryEvent.objects.create(delivery=delivery,actor=request.user,action='reordered')
    return redirect('courier_today')
